"""
CV text extraction (PDF / DOCX).

PROVENANCE NOTE: `src/cv/parser.py` referenced by the notebook (with a
documented "never persist uploaded documents" policy) was not
uploaded. This is a standard, from-scratch text-extraction
implementation - not a reconstruction of proprietary logic, since
extracting text from a PDF/DOCX is a well-defined mechanical task with
no hidden model behavior to guess at. The "never persist the raw
document" policy mentioned in the notebook IS honored here: callers
should extract text and discard the original bytes/path once done.
"""

from __future__ import annotations

from pathlib import Path

SUPPORTED_EXTENSIONS = {".pdf", ".docx"}


class UnsupportedFileTypeError(ValueError):
    pass


class CVParsingError(RuntimeError):
    pass


def extract_text(file_path: str | Path) -> str:
    path = Path(file_path)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _extract_pdf(path)
    if suffix == ".docx":
        return _extract_docx(path)
    raise UnsupportedFileTypeError(f"Unsupported CV file type: {suffix}")


def _extract_pdf(path: Path) -> str:
    try:
        import pdfplumber
    except ImportError as exc:
        raise CVParsingError("pdfplumber is required to parse PDF CVs") from exc

    try:
        text_parts = []
        with pdfplumber.open(path) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text() or ""
                text_parts.append(page_text)
        text = "\n".join(text_parts).strip()
        if not text:
            raise CVParsingError("No extractable text found in PDF (likely a scanned image).")
        return text
    except CVParsingError:
        raise
    except Exception as exc:
        raise CVParsingError(f"Failed to parse PDF: {exc}") from exc


def _extract_docx(path: Path) -> str:
    try:
        import docx
    except ImportError as exc:
        raise CVParsingError("python-docx is required to parse DOCX CVs") from exc

    try:
        document = docx.Document(str(path))
        paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        paragraphs.append(cell.text)
        text = "\n".join(paragraphs).strip()
        if not text:
            raise CVParsingError("No extractable text found in DOCX.")
        return text
    except CVParsingError:
        raise
    except Exception as exc:
        raise CVParsingError(f"Failed to parse DOCX: {exc}") from exc


def validate_file(file_path: str | Path, max_size_bytes: int = 10 * 1024 * 1024) -> None:
    path = Path(file_path)
    if not path.is_file():
        raise CVParsingError("Uploaded file could not be found on disk.")
    if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFileTypeError(f"Unsupported file type: {path.suffix}")
    size = path.stat().st_size
    if size == 0:
        raise CVParsingError("Uploaded file is empty.")
    if size > max_size_bytes:
        raise CVParsingError(f"File exceeds maximum size of {max_size_bytes} bytes.")

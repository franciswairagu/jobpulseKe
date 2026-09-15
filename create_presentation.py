#!/usr/bin/env python3
"""Generate JobPulse presentation — accessible, story-driven, max 13 slides."""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# ── Palette ─────────────────────────────────────────────────────────
BG        = RGBColor(0x0F, 0x17, 0x2A)
CARD      = RGBColor(0x16, 0x21, 0x3E)
CYAN      = RGBColor(0x00, 0xD2, 0xFF)
PURPLE    = RGBColor(0x7C, 0x3A, 0xED)
GREEN     = RGBColor(0x10, 0xB9, 0x81)
AMBER     = RGBColor(0xF5, 0x9E, 0x0B)
RED       = RGBColor(0xEF, 0x44, 0x44)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
LGRAY     = RGBColor(0xCB, 0xD5, 0xE1)
MGRAY     = RGBColor(0x94, 0xA3, 0xB8)

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)


# ── Helpers ─────────────────────────────────────────────────────────
def bg(s):
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG

def box(s, l, t, w, h, fill=None, border=None, bw=Pt(1)):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    sh.shadow.inherit = False
    if fill:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
    else:
        sh.fill.background()
    if border:
        sh.line.color.rgb = border; sh.line.width = bw
    else:
        sh.line.fill.background()
    return sh

def rect(s, l, t, w, h, fill):
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    sh.shadow.inherit = False
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    return sh

def txt(s, l, t, w, h, text, sz=14, clr=WHITE, bold=False, align=PP_ALIGN.LEFT):
    tb = s.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text; p.font.size = Pt(sz); p.font.color.rgb = clr
    p.font.bold = bold; p.font.name = "Calibri"; p.alignment = align
    return tb

def circle(s, l, t, d, clr):
    sh = s.shapes.add_shape(MSO_SHAPE.OVAL, l, t, d, d)
    sh.fill.solid(); sh.fill.fore_color.rgb = clr
    sh.line.fill.background(); sh.shadow.inherit = False
    return sh

def accent(s, l, t, w, clr=CYAN):
    rect(s, l, t, w, Pt(4), clr)

def title_bar(s, heading, sub=None):
    txt(s, Inches(0.8), Inches(0.4), Inches(11), Inches(0.7), heading, sz=30, clr=WHITE, bold=True)
    accent(s, Inches(0.8), Inches(1.05), Inches(2.5))
    if sub:
        txt(s, Inches(0.8), Inches(1.25), Inches(11), Inches(0.5), sub, sz=14, clr=MGRAY)


# ═══════════════════════════════════════════════════════════════════
# 1 — Title
# ═══════════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(prs.slide_layouts[6]); bg(sl)

circle(sl, Inches(9.2), Inches(-1.5), Inches(9), RGBColor(0x0D, 0x2A, 0x3E))

txt(sl, Inches(0.8), Inches(1.8), Inches(8), Inches(1.0),
    "JobPulse", sz=56, clr=CYAN, bold=True)
accent(sl, Inches(0.8), Inches(2.85), Inches(4))
txt(sl, Inches(0.8), Inches(3.2), Inches(9), Inches(1.0),
    "Helping African tech talent find the right job\n— and close the skills gap to get there.",
    sz=22, clr=WHITE)
txt(sl, Inches(0.8), Inches(4.6), Inches(9), Inches(0.5),
    "Moringa School  ·  DSF-FT16 Capstone  ·  September 2026",
    sz=13, clr=MGRAY)

box(sl, Inches(0.8), Inches(6.0), Inches(11.7), Inches(0.9), fill=CARD)
txt(sl, Inches(1.1), Inches(6.1), Inches(11), Inches(0.7),
    "10,379 real job postings  ·  15+ sources  ·  10 African countries  ·  600+ skills tracked",
    sz=14, clr=LGRAY, align=PP_ALIGN.CENTER)


# ═══════════════════════════════════════════════════════════════════
# 2 — The Problem (the story)
# ═══════════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(prs.slide_layouts[6]); bg(sl)
title_bar(sl, "The Problem", "This isn't hypothetical — it's happening right now")

box(sl, Inches(0.8), Inches(1.8), Inches(11.7), Inches(4.0), fill=CARD, border=CYAN, bw=Pt(2))

story = (
    "Picture a computer science graduate in Nairobi.\n\n"
    "She's applied to forty jobs in three months. She spends an evening "
    "tailoring her CV for one role, only to find out two weeks later — "
    "after finally hearing back — that the position was filled before "
    "she even applied.\n\n"
    "The listing was still live. Nobody took it down."
)
txt(sl, Inches(1.3), Inches(2.0), Inches(10.7), Inches(3.5), story, sz=17, clr=LGRAY)

txt(sl, Inches(1.3), Inches(5.4), Inches(10.7), Inches(0.5),
    "And when she doesn't hear back at all, she has no idea why. "
    "Nobody tells her what stands between her and the next interview.",
    sz=15, clr=MGRAY)


# ═══════════════════════════════════════════════════════════════════
# 3 — Why This Happens
# ═══════════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(prs.slide_layouts[6]); bg(sl)
title_bar(sl, "Why This Happens", "The job market is broken in ways nobody talks about")

problems = [
    ("Dead Listings", "Job posts stay live long after the role\nis filled — wasting everyone's time.", RED),
    ("Zero Feedback", "Rejections come with no explanation.\nCandidates can't learn what went wrong.", AMBER),
    ("No Skill Visibility", "You don't know which skills the market\nactually wants right now.", PURPLE),
    ("Fragmented Data", "Job info is scattered across dozens of\nsites — nobody sees the full picture.", CYAN),
]
for i, (title, desc, clr) in enumerate(problems):
    x = Inches(0.6) + i * Inches(3.15)
    box(sl, x, Inches(1.8), Inches(2.9), Inches(4.2), fill=CARD, border=clr, bw=Pt(1.5))
    circle(sl, x + Inches(1.05), Inches(2.1), Inches(0.7), clr)
    txt(sl, x + Inches(1.05), Inches(2.15), Inches(0.7), Inches(0.6),
        str(i+1), sz=22, clr=BG, bold=True, align=PP_ALIGN.CENTER)
    txt(sl, x + Inches(0.2), Inches(3.0), Inches(2.5), Inches(0.5),
        title, sz=16, clr=WHITE, bold=True, align=PP_ALIGN.CENTER)
    txt(sl, x + Inches(0.2), Inches(3.6), Inches(2.5), Inches(2.0),
        desc, sz=13, clr=LGRAY, align=PP_ALIGN.CENTER)

txt(sl, Inches(0.8), Inches(6.4), Inches(11.7), Inches(0.5),
    "The result? Talented people give up. Employers miss great candidates. Nobody wins.",
    sz=14, clr=MGRAY, align=PP_ALIGN.CENTER)


# ═══════════════════════════════════════════════════════════════════
# 4 — Someone Tried Before
# ═══════════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(prs.slide_layouts[6]); bg(sl)
title_bar(sl, "Someone Tried Before", "A promising experiment that never became a product")

box(sl, Inches(0.8), Inches(1.8), Inches(5.5), Inches(4.8), fill=CARD, border=CYAN, bw=Pt(1))
txt(sl, Inches(1.1), Inches(1.95), Inches(5), Inches(0.5),
    "World Bank × Headai (2019)", sz=18, clr=CYAN, bold=True)

story2 = (
    "The World Bank partnered with a Finnish AI company to crawl over "
    "60,000 job postings from three Kenyan job boards.\n\n"
    "They compared what employers were asking for against what universities "
    "were teaching.\n\n"
    "It worked. They found a real gap between skills taught and skills demanded."
)
txt(sl, Inches(1.1), Inches(2.6), Inches(5.0), Inches(3.5), story2, sz=14, clr=LGRAY)

box(sl, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.8), fill=CARD, border=RED, bw=Pt(1))
txt(sl, Inches(7.1), Inches(1.95), Inches(5), Inches(0.5),
    "Then It Stopped", sz=18, clr=RED, bold=True)

ended = (
    "It was a two-month pilot on a fixed dataset from 2015 to 2019.\n\n"
    "Never repeated. Never kept current. Never turned into something "
    "people could actually use.\n\n"
    "The approach was validated.\nNobody had turned it into a product."
)
txt(sl, Inches(7.1), Inches(2.6), Inches(5.0), Inches(3.5), ended, sz=14, clr=LGRAY)

# Bottom
box(sl, Inches(0.8), Inches(6.8), Inches(11.7), Inches(0.55),
    fill=RGBColor(0x0A, 0x2A, 0x1A), border=GREEN, bw=Pt(1))
txt(sl, Inches(1.1), Inches(6.83), Inches(11), Inches(0.5),
    "✦  That's the gap JobPulse is built to close.",
    sz=15, clr=GREEN, bold=True, align=PP_ALIGN.CENTER)


# ═══════════════════════════════════════════════════════════════════
# 5 — Introducing JobPulse
# ═══════════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(prs.slide_layouts[6]); bg(sl)
title_bar(sl, "Introducing JobPulse", "A complete career companion — not just another job board")

txt(sl, Inches(0.8), Inches(1.6), Inches(11.7), Inches(0.8),
    "JobPulse collects real job data from across Africa, analyzes it, and gives job seekers "
    "personalized tools to understand the market, improve their CVs, and find the right roles.",
    sz=16, clr=LGRAY)

pillars = [
    ("Know Your\nStrengths", "Upload your CV and get\na clear picture of what\nyou offer — and what's\nmissing.", CYAN),
    ("Find the\nRight Jobs", "See roles that actually\nmatch your skills, ranked\nby how well you fit —\nnot just keywords.", PURPLE),
    ("Skip Dead\nListings", "We track which jobs are\nstill open, filled, or\nexpired — so you don't\nwaste your time.", GREEN),
    ("Get\nGuidance", "Ask our AI assistant\nanything — it points you\nto courses, practice\nresources, and next steps.", AMBER),
]
for i, (title, desc, clr) in enumerate(pillars):
    x = Inches(0.6) + i * Inches(3.15)
    box(sl, x, Inches(2.8), Inches(2.9), Inches(3.6), fill=CARD, border=clr, bw=Pt(1.5))
    circle(sl, x + Inches(1.1), Inches(3.0), Inches(0.6), clr)
    txt(sl, x + Inches(1.1), Inches(3.05), Inches(0.6), Inches(0.5),
        str(i+1), sz=20, clr=BG, bold=True, align=PP_ALIGN.CENTER)
    txt(sl, x + Inches(0.2), Inches(3.8), Inches(2.5), Inches(0.8),
        title, sz=15, clr=WHITE, bold=True, align=PP_ALIGN.CENTER)
    txt(sl, x + Inches(0.2), Inches(4.7), Inches(2.5), Inches(1.5),
        desc, sz=12, clr=LGRAY, align=PP_ALIGN.CENTER)


# ═══════════════════════════════════════════════════════════════════
# 6 — Feature: CV Analysis
# ═══════════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(prs.slide_layouts[6]); bg(sl)
title_bar(sl, "Your CV, Analyzed", "Upload it — we'll tell you exactly where you stand")

steps = [
    ("Upload", "Drop your PDF,\nWord doc, or text file"),
    ("We Read It", "Our system extracts\nyour skills, experience,\neducation, and certs"),
    ("Get Your Score", "A clear 0–100 score\nshowing how your CV\nstacks up to the market"),
    ("See Your Gaps", "We compare you against\nwhat employers actually\nwant right now"),
    ("Take Action", "Get matched to jobs\nand courses that close\nyour specific gaps"),
]
for i, (t, d) in enumerate(steps):
    x = Inches(0.4) + i * Inches(2.5)
    box(sl, x, Inches(1.8), Inches(2.2), Inches(2.8), fill=CARD, border=CYAN, bw=Pt(1))
    txt(sl, x, Inches(1.95), Inches(2.2), Inches(0.5),
        f"0{i+1}", sz=24, clr=CYAN, bold=True, align=PP_ALIGN.CENTER)
    txt(sl, x, Inches(2.5), Inches(2.2), Inches(0.5),
        t, sz=15, clr=WHITE, bold=True, align=PP_ALIGN.CENTER)
    txt(sl, x + Inches(0.15), Inches(3.1), Inches(1.9), Inches(1.2),
        d, sz=12, clr=LGRAY, align=PP_ALIGN.CENTER)
    if i < 4:
        txt(sl, x + Inches(2.2), Inches(2.7), Inches(0.3), Inches(0.5),
            "→", sz=20, clr=CYAN, bold=True, align=PP_ALIGN.CENTER)

# What you see
box(sl, Inches(0.8), Inches(5.0), Inches(11.7), Inches(2.2), fill=CARD)
txt(sl, Inches(1.1), Inches(5.1), Inches(5), Inches(0.5),
    "What You Get Back", sz=16, clr=CYAN, bold=True)

items = [
    "Your extracted skills organized by category (programming, cloud, databases, etc.)",
    "A breakdown of your score: skills, experience, education, and certifications",
    "A list of missing skills ranked by how often they appear in job postings",
    "Specific job listings where you're a strong match — and where you fall short",
]
for i, item in enumerate(items):
    y = Inches(5.7) + i * Inches(0.38)
    txt(sl, Inches(1.3), y, Inches(10.5), Inches(0.35),
        f"▸  {item}", sz=12, clr=LGRAY)


# ═══════════════════════════════════════════════════════════════════
# 7 — Feature: Smart Job Matching
# ═══════════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(prs.slide_layouts[6]); bg(sl)
title_bar(sl, "Smart Job Matching", "Not just keyword search — real skill-based matching")

box(sl, Inches(0.8), Inches(1.8), Inches(6.0), Inches(5.0), fill=CARD, border=PURPLE, bw=Pt(1))
txt(sl, Inches(1.1), Inches(1.9), Inches(5.5), Inches(0.5),
    "How We Rank Your Matches", sz=18, clr=PURPLE, bold=True)

factors = [
    ("Do you have the skills they need?", "65%", "This is the biggest factor — the core skills the role requires."),
    ("Do you have the nice-to-haves?", "15%", "Preferred skills that make you stand out from other candidates."),
    ("Do you have enough experience?", "15%", "Your years of experience compared to what the role asks for."),
    ("Does the location work?", "5%", "Country match and whether you prefer remote, hybrid, or onsite."),
]
for i, (q, w, d) in enumerate(factors):
    y = Inches(2.5) + i * Inches(1.0)
    txt(sl, Inches(1.3), y, Inches(4.5), Inches(0.35), q, sz=13, clr=WHITE, bold=True)
    txt(sl, Inches(5.8), y, Inches(0.8), Inches(0.35), w, sz=16, clr=PURPLE, bold=True, align=PP_ALIGN.CENTER)
    txt(sl, Inches(1.3), y + Inches(0.35), Inches(5.0), Inches(0.4), d, sz=11, clr=MGRAY)

box(sl, Inches(7.2), Inches(1.8), Inches(5.3), Inches(5.0), fill=CARD, border=GREEN, bw=Pt(1))
txt(sl, Inches(7.5), Inches(1.9), Inches(4.7), Inches(0.5),
    "And Then We Build You a Plan", sz=18, clr=GREEN, bold=True)

plan_text = (
    "Once we know which jobs fit you best, we look at what skills "
    "you're missing across all those matches.\n\n"
    "We prioritize the skills that come up most often — the ones that "
    "would open the most doors for you.\n\n"
    "Then we find courses, tutorials, and practice resources for each "
    "one — so you know exactly what to work on and where to go learn it."
)
txt(sl, Inches(7.5), Inches(2.6), Inches(4.7), Inches(4.0), plan_text, sz=13, clr=LGRAY)


# ═══════════════════════════════════════════════════════════════════
# 8 — Feature: Job Status Tracking
# ═══════════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(prs.slide_layouts[6]); bg(sl)
title_bar(sl, "No More Wasted Applications", "We track which jobs are real — so you don't apply to ghosts")

statuses = [
    ("Still Open", GREEN, "This job is accepting\napplications right now.\nGo for it."),
    ("Filled / Expired", RED, "This role is no longer\navailable. Don't waste\nyour evening on it."),
    ("Unknown", AMBER, "We haven't checked\nthis one yet. We're\nworking on it."),
]
for i, (label, clr, desc) in enumerate(statuses):
    x = Inches(0.8) + i * Inches(4.0)
    box(sl, x, Inches(1.8), Inches(3.6), Inches(2.5), fill=CARD, border=clr, bw=Pt(1.5))
    txt(sl, x, Inches(2.0), Inches(3.6), Inches(0.5),
        label, sz=20, clr=clr, bold=True, align=PP_ALIGN.CENTER)
    txt(sl, x, Inches(2.7), Inches(3.6), Inches(1.5),
        desc, sz=14, clr=LGRAY, align=PP_ALIGN.CENTER)

# How it works — plain language
box(sl, Inches(0.8), Inches(4.7), Inches(11.7), Inches(2.5), fill=CARD)
txt(sl, Inches(1.1), Inches(4.8), Inches(5), Inches(0.5),
    "How We Keep Track", sz=16, clr=AMBER, bold=True)

explain = [
    "When we first find a job, we mark it as available.",
    "We periodically re-check listings — if a job disappears from the site, we update the status.",
    "If a job hasn't been seen in 30 days, we flag it as likely removed.",
    "Every change is logged, so there's a clear history of what happened and when.",
]
for i, item in enumerate(explain):
    txt(sl, Inches(1.3), Inches(5.4) + i * Inches(0.42), Inches(10.5), Inches(0.4),
        f"▸  {item}", sz=13, clr=LGRAY)


# ═══════════════════════════════════════════════════════════════════
# 9 — Feature: AI Assistant
# ═══════════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(prs.slide_layouts[6]); bg(sl)
title_bar(sl, "Your AI Career Assistant", "Just ask — it knows the market and it knows your profile")

# Example questions
box(sl, Inches(0.8), Inches(1.8), Inches(5.5), Inches(5.0), fill=CARD, border=AMBER, bw=Pt(1))
txt(sl, Inches(1.1), Inches(1.9), Inches(5), Inches(0.5),
    "Things You Can Ask", sz=18, clr=AMBER, bold=True)

examples = [
    ("\"What skills do most DevOps roles in Nairobi require?\"",
     "Market intelligence — real demand data."),
    ("\"I know Python and SQL. What jobs match me?\"",
     "Job search — matched to your actual skills."),
    ("\"What's the average salary for a data analyst in Lagos?\"",
     "Salary insights — from real postings."),
    ("\"I want to move into cloud engineering. Where do I start?\"",
     "Career advice — personalized learning path."),
    ("\"Is there high demand for React developers in South Africa?\"",
     "Regional analysis — skill demand by country."),
]
for i, (q, a) in enumerate(examples):
    y = Inches(2.5) + i * Inches(0.85)
    txt(sl, Inches(1.3), y, Inches(5.0), Inches(0.35), q, sz=12, clr=WHITE, bold=True)
    txt(sl, Inches(1.3), y + Inches(0.35), Inches(5.0), Inches(0.35), a, sz=11, clr=MGRAY)

# How it works — simple
box(sl, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.0), fill=CARD, border=CYAN, bw=Pt(1))
txt(sl, Inches(7.1), Inches(1.9), Inches(5), Inches(0.5),
    "How It Works (Simply)", sz=18, clr=CYAN, bold=True)

simple_arch = (
    "1.  You ask a question in plain language.\n\n"
    "2.  We search our database of real job postings to find "
    "the most relevant ones.\n\n"
    "3.  Our AI reads those results and puts together a clear, "
    "helpful answer — with specific examples and links.\n\n"
    "4.  Every answer is grounded in real data — no made-up "
    "information, no generic advice."
)
txt(sl, Inches(7.1), Inches(2.6), Inches(5.0), Inches(4.0), simple_arch, sz=13, clr=LGRAY)

# Stats in plain language
txt(sl, Inches(7.1), Inches(5.8), Inches(5.0), Inches(0.8),
    "Accuracy: 97% of the time, the first result it finds is exactly what you need. "
    "100% of answers are backed by real job data.",
    sz=12, clr=GREEN)


# ═══════════════════════════════════════════════════════════════════
# 10 — Real Data, Real Impact
# ═══════════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(prs.slide_layouts[6]); bg(sl)
title_bar(sl, "Real Data, Real Impact", "What we've collected and what it tells us")

stats = [
    ("10,379", "Job postings\nscraped and analyzed", CYAN),
    ("15+", "Sources across\nAfrica and beyond", PURPLE),
    ("10", "African countries\nrepresented", GREEN),
    ("720+", "Unique companies\nhiring", AMBER),
    ("600+", "Distinct skills\ntracked", RED),
]
for i, (num, label, clr) in enumerate(stats):
    x = Inches(0.5) + i * Inches(2.5)
    box(sl, x, Inches(1.8), Inches(2.2), Inches(2.5), fill=CARD, border=clr, bw=Pt(1.5))
    txt(sl, x, Inches(2.0), Inches(2.2), Inches(0.8),
        num, sz=34, clr=clr, bold=True, align=PP_ALIGN.CENTER)
    txt(sl, x, Inches(2.9), Inches(2.2), Inches(1.0),
        label, sz=13, clr=LGRAY, align=PP_ALIGN.CENTER)

# What this enables
box(sl, Inches(0.8), Inches(4.7), Inches(11.7), Inches(2.5), fill=CARD)
txt(sl, Inches(1.1), Inches(4.8), Inches(5), Inches(0.5),
    "What This Data Unlocks", sz=16, clr=CYAN, bold=True)

unlocks = [
    "See which skills are most in-demand across different African countries",
    "Understand how requirements differ between senior and junior roles",
    "Track which job boards have the most active listings in your region",
    "Compare remote vs. onsite opportunities by skill set and location",
    "Identify emerging skill trends before they become mainstream",
]
for i, item in enumerate(unlocks):
    txt(sl, Inches(1.3), Inches(5.4) + i * Inches(0.38), Inches(10.5), Inches(0.35),
        f"▸  {item}", sz=13, clr=LGRAY)


# ═══════════════════════════════════════════════════════════════════
# 11 — How It All Fits Together
# ═══════════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(prs.slide_layouts[6]); bg(sl)
title_bar(sl, "How It All Fits Together", "From raw data to career action")

flow = [
    ("Scrape", "Collect jobs from\n15+ websites", CYAN),
    ("Clean", "Remove duplicates,\nfix locations, standardize", PURPLE),
    ("Analyze", "Extract skills,\nexperience, education", GREEN),
    ("Store", " organized database\nready for querying", AMBER),
    ("Serve", "CV analysis, job\nmatching, AI assistant", RED),
]
for i, (t, d, clr) in enumerate(flow):
    x = Inches(0.4) + i * Inches(2.5)
    box(sl, x, Inches(1.8), Inches(2.2), Inches(2.5), fill=CARD, border=clr, bw=Pt(1.5))
    circle(sl, x + Inches(0.75), Inches(2.0), Inches(0.6), clr)
    txt(sl, x + Inches(0.75), Inches(2.05), Inches(0.6), Inches(0.5),
        str(i+1), sz=20, clr=BG, bold=True, align=PP_ALIGN.CENTER)
    txt(sl, x, Inches(2.8), Inches(2.2), Inches(0.4),
        t, sz=15, clr=WHITE, bold=True, align=PP_ALIGN.CENTER)
    txt(sl, x + Inches(0.15), Inches(3.25), Inches(1.9), Inches(0.8),
        d, sz=12, clr=LGRAY, align=PP_ALIGN.CENTER)
    if i < 4:
        txt(sl, x + Inches(2.2), Inches(2.6), Inches(0.3), Inches(0.5),
            "→", sz=20, clr=LIGHT_GRAY if 'LIGHT_GRAY' in dir() else LGRAY, bold=True, align=PP_ALIGN.CENTER)

# User journey
box(sl, Inches(0.8), Inches(4.7), Inches(11.7), Inches(2.5), fill=CARD)
txt(sl, Inches(1.1), Inches(4.8), Inches(5), Inches(0.5),
    "Your Journey", sz=16, clr=GREEN, bold=True)

journey = [
    "Sign up for free →",
    "Upload your CV →",
    "Get your skills analysis and score →",
    "See matched jobs ranked by fit →",
    "Get a personalized learning plan →",
    "Ask the AI assistant for guidance anytime",
]
for i, item in enumerate(journey):
    txt(sl, Inches(1.3), Inches(5.4) + i * Inches(0.35), Inches(10.5), Inches(0.35),
        f"  {i+1}.  {item}", sz=13, clr=LGRAY)


# ═══════════════════════════════════════════════════════════════════
# 12 — What's Next
# ═══════════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(prs.slide_layouts[6]); bg(sl)
title_bar(sl, "What's Next", "The roadmap ahead")

box(sl, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0), fill=CARD, border=PURPLE, bw=Pt(1))
txt(sl, Inches(1.1), Inches(1.9), Inches(5), Inches(0.5),
    "Short Term", sz=18, clr=PURPLE, bold=True)
short = [
    "Automated daily re-scraping to keep job data fresh",
    "Improved skill detection for emerging technologies",
    "Conversation memory — remember your previous questions",
    "Mobile-friendly interface for access anywhere",
]
for i, item in enumerate(short):
    txt(sl, Inches(1.3), Inches(2.6) + i * Inches(0.5), Inches(5.0), Inches(0.45),
        f"▸  {item}", sz=13, clr=LGRAY)

box(sl, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.0), fill=CARD, border=GREEN, bw=Pt(1))
txt(sl, Inches(7.1), Inches(1.9), Inches(5), Inches(0.5),
    "Long Term", sz=18, clr=GREEN, bold=True)
long = [
    "Expand to more African countries and job boards",
    "Employer dashboard — help companies understand their hiring market",
    "Community features — connect job seekers with mentors",
    "Integration with learning platforms for seamless course enrollment",
    "Scale to handle 100,000+ job postings",
]
for i, item in enumerate(long):
    txt(sl, Inches(7.3), Inches(2.6) + i * Inches(0.5), Inches(5.0), Inches(0.45),
        f"▸  {item}", sz=13, clr=LGRAY)


# ═══════════════════════════════════════════════════════════════════
# 13 — Closing
# ═══════════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(prs.slide_layouts[6]); bg(sl)

txt(sl, Inches(0.8), Inches(1.5), Inches(11.7), Inches(1.0),
    "JobPulse", sz=52, clr=CYAN, bold=True, align=PP_ALIGN.CENTER)
accent(sl, Inches(5.4), Inches(2.5), Inches(2.5))
txt(sl, Inches(1.5), Inches(2.9), Inches(10.3), Inches(1.0),
    "Because no one should wonder why\nthey didn't get the interview.",
    sz=24, clr=WHITE, align=PP_ALIGN.CENTER)

# Three columns
cols = [
    ("For Job Seekers", CYAN, [
        "Know your strengths and gaps",
        "Find jobs that actually match you",
        "Skip dead listings",
        "Get personalized learning plans",
        "Ask an AI assistant anything",
    ]),
    ("For the Market", GREEN, [
        "Real-time skill demand data",
        "Track which roles are actually open",
        "Understand regional differences",
        "Identify emerging trends",
        "Bridge the education gap",
    ]),
    ("Built for Africa", AMBER, [
        "10+ African countries covered",
        "Local job boards included",
        "Region-specific insights",
        "Works offline with local AI",
        "Privacy-first design",
    ]),
]
for i, (title, clr, items) in enumerate(cols):
    x = Inches(0.8) + i * Inches(4.1)
    box(sl, x, Inches(4.2), Inches(3.8), Inches(2.8), fill=CARD, border=clr, bw=Pt(1))
    txt(sl, x + Inches(0.2), Inches(4.3), Inches(3.4), Inches(0.5),
        title, sz=15, clr=clr, bold=True, align=PP_ALIGN.CENTER)
    for j, item in enumerate(items):
        txt(sl, x + Inches(0.3), Inches(4.85) + j * Inches(0.38), Inches(3.2), Inches(0.35),
            f"▸  {item}", sz=11, clr=LGRAY)


# ── Save ────────────────────────────────────────────────────────────
out = "/home/wairagu/Desktop/Moringa/DSF-FT16/module 6/jobpulse/JobPulse_Presentation.pptx"
prs.save(out)
print(f"Saved: {out}")
